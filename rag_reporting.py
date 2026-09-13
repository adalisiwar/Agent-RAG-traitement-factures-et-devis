import base64
import os
from datetime import datetime

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from db_utils import execute_with_retry
from rag_search import BASE_FROM, DATABASE_URL, build_where, detect_filters

NAVY = "1E2761"
MIDBLUE = "3B5BA5"
ICE = "CADCFC"
MUTED = "5A6B87"
LIGHT_BG = "F4F6FB"

REPORT_KEYWORDS = [
    "rapport", "tableau", "graphique", "visualisation", "visualiser",
    "répartition", "repartition", "évolution", "evolution",
]

GROUP_BY_OPTIONS = {
    "fournisseur": ("f.nom", "Fournisseur"),
    "statut": ("s.libelle", "Statut"),
    "type": ("d.type_document", "Type de document"),
    "mois": ("to_char(d.date_emission, 'YYYY-MM')", "Mois"),
    "annee": ("EXTRACT(YEAR FROM d.date_emission)::text", "Année"),
    "devise": ("d.devise", "Devise"),
}

METRIC_COLUMN_BY_LABEL = {
    "Montant TTC total": "Montant TTC",
    "Reste à payer total": "Reste à payer",
    "Montant payé total": "Montant payé",
    "Nombre de documents": "Nombre de documents",
}

METRIC_LABELS = {
    "montant_ttc": "Montant TTC total",
    "reste_a_payer": "Reste à payer total",
    "montant_paye": "Montant payé total",
    "count": "Nombre de documents",
}

TOP_N_IN_CHART = 8


def is_report_request(question):
    question_lower = question.lower()
    return any(keyword in question_lower for keyword in REPORT_KEYWORDS)


def detect_group_by(question):
    question_lower = question.lower()
    for key, (expr, label) in GROUP_BY_OPTIONS.items():
        if key in question_lower:
            return expr, label
    return GROUP_BY_OPTIONS["fournisseur"]


def detect_metric(question):
    question_lower = question.lower()
    if any(w in question_lower for w in ["impayé", "impaye", "reste à payer", "reste a payer"]):
        return "reste_a_payer"
    if any(w in question_lower for w in ["nombre", "combien"]):
        return "count"
    if "payé" in question_lower or "paye" in question_lower:
        return "montant_paye"
    return "montant_ttc"


def build_report_dataframe(conn_holder, question):
    filters = detect_filters(question)
    where_sql, params = build_where(filters)

    group_expr, group_label = detect_group_by(question)
    metric_key = detect_metric(question)
    metric_label = METRIC_LABELS[metric_key]

    sql = f"""
    SELECT
        {group_expr} AS groupe,
        COUNT(*) AS nb_documents,
        COALESCE(SUM(d.montant_ttc), 0) AS montant_ttc,
        COALESCE(SUM(d.montant_paye), 0) AS montant_paye,
        COALESCE(SUM(d.reste_a_payer), 0) AS reste_a_payer
    {BASE_FROM}
    {where_sql}
    GROUP BY {group_expr}
    ORDER BY 3 DESC
    LIMIT 20
    """

    rows = execute_with_retry(DATABASE_URL, conn_holder, sql, params)

    df = pd.DataFrame(rows, columns=[
        group_label, "Nombre de documents", "Montant TTC",
        "Montant payé", "Reste à payer",
    ])

    numeric_columns = ["Nombre de documents", "Montant TTC", "Montant payé", "Reste à payer"]
    df[numeric_columns] = df[numeric_columns].astype(float)

    return df, group_label, metric_label


def _hex_to_rgb(hex_color):
    hex_color = hex_color.lstrip("#")
    return tuple(int(hex_color[i:i + 2], 16) / 255 for i in (0, 2, 4))


def _gradient(n, start_hex=NAVY, end_hex=ICE):
    start = _hex_to_rgb(start_hex)
    end = _hex_to_rgb(end_hex)
    if n <= 1:
        return [f"#{start_hex}"]
    colors = []
    for i in range(n):
        t = i / (n - 1)
        rgb = tuple(start[c] + (end[c] - start[c]) * t for c in range(3))
        colors.append("#{:02x}{:02x}{:02x}".format(
            int(rgb[0] * 255), int(rgb[1] * 255), int(rgb[2] * 255)
        ))
    return colors


def _format_number(value):
    return f"{value:,.0f}".replace(",", " ")


def _prepare_chart_data(df, group_label, metric_column):
    working = df[[group_label, metric_column]].copy()
    working = working.sort_values(metric_column, ascending=False).reset_index(drop=True)

    if len(working) > TOP_N_IN_CHART:
        top = working.iloc[:TOP_N_IN_CHART].copy()
        others_total = working.iloc[TOP_N_IN_CHART:][metric_column].sum()
        others_row = pd.DataFrame({group_label: ["Autres"], metric_column: [others_total]})
        working = pd.concat([top, others_row], ignore_index=True)

    return working.iloc[::-1].reset_index(drop=True)


def render_chart(df, group_label, metric_label, output_dir, filename="graphique.png"):
    os.makedirs(output_dir, exist_ok=True)
    path = os.path.join(output_dir, filename)

    metric_column = METRIC_COLUMN_BY_LABEL.get(metric_label, "Montant TTC")
    chart_df = _prepare_chart_data(df, group_label, metric_column)

    colors = _gradient(len(chart_df))
    colors = colors[::-1]

    fig, ax = plt.subplots(figsize=(11, max(4.5, 0.55 * len(chart_df))))
    bars = ax.barh(chart_df[group_label].astype(str), chart_df[metric_column], color=colors, height=0.62)

    for bar, value in zip(bars, chart_df[metric_column]):
        ax.text(
            bar.get_width() + max(chart_df[metric_column]) * 0.01,
            bar.get_y() + bar.get_height() / 2,
            _format_number(value),
            va="center", ha="left", fontsize=10, color=f"#{NAVY}",
        )

    ax.set_xlabel(metric_column, fontsize=11, color=f"#{MUTED}")
    ax.set_title(f"{metric_column} par {group_label}", fontsize=15, color=f"#{NAVY}", weight="bold", pad=16, loc="left")

    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.spines["left"].set_visible(False)
    ax.spines["bottom"].set_color(f"#{ICE}")
    ax.tick_params(axis="y", length=0, labelsize=10.5, colors=f"#{NAVY}")
    ax.tick_params(axis="x", length=0, labelsize=9.5, colors=f"#{MUTED}")
    ax.xaxis.grid(True, color=f"#{ICE}", linewidth=0.8, zorder=0)
    ax.set_axisbelow(True)
    ax.margins(x=0.12)

    fig.patch.set_facecolor("white")
    ax.set_facecolor("white")
    fig.tight_layout()
    fig.savefig(path, dpi=170, facecolor="white")
    plt.close(fig)

    return path


def export_excel(df, output_dir, filename="rapport.xlsx"):
    os.makedirs(output_dir, exist_ok=True)
    path = os.path.join(output_dir, filename)

    numeric_columns = ["Montant TTC", "Montant payé", "Reste à payer"]
    total_row = {col: df[col].sum() if col in numeric_columns else "" for col in df.columns}
    total_row[df.columns[0]] = "Total"
    export_df = pd.concat([df, pd.DataFrame([total_row])], ignore_index=True)

    export_df.to_excel(path, index=False, engine="openpyxl", sheet_name="Rapport")

    from openpyxl import load_workbook
    workbook = load_workbook(path)
    sheet = workbook["Rapport"]

    header_fill = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    total_fill = PatternFill(start_color=ICE, end_color=ICE, fill_type="solid")
    total_font = Font(color=NAVY, bold=True)
    thin_border = Border(bottom=Side(style="thin", color=ICE))

    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    last_row = sheet.max_row
    for cell in sheet[last_row]:
        cell.fill = total_fill
        cell.font = total_font

    for row in sheet.iter_rows(min_row=2, max_row=last_row - 1):
        for cell in row:
            cell.border = thin_border

    for col_index, column in enumerate(export_df.columns, start=1):
        letter = get_column_letter(col_index)
        max_length = max(
            [len(str(column))] + [len(str(v)) for v in export_df[column].tolist()]
        )
        sheet.column_dimensions[letter].width = max_length + 4

        if column in numeric_columns:
            for row_index in range(2, last_row + 1):
                sheet.cell(row=row_index, column=col_index).number_format = "#,##0.00"

    sheet.freeze_panes = "A2"
    workbook.save(path)

    return path


def _image_to_base64(path):
    with open(path, "rb") as file:
        return base64.b64encode(file.read()).decode("utf-8")


def _render_table_html(df):
    numeric_columns = ["Montant TTC", "Montant payé", "Reste à payer"]

    header_cells = "".join(f"<th>{col}</th>" for col in df.columns)
    body_rows = []
    for i, row in df.iterrows():
        row_class = "alt" if i % 2 else ""
        cells = []
        for col in df.columns:
            value = row[col]
            if col in numeric_columns:
                cells.append(f"<td class='num'>{_format_number(value)}</td>")
            elif col == "Nombre de documents":
                cells.append(f"<td class='num'>{value}</td>")
            else:
                cells.append(f"<td>{value}</td>")
        body_rows.append(f"<tr class='{row_class}'>{''.join(cells)}</tr>")

    return f"<table><thead><tr>{header_cells}</tr></thead><tbody>{''.join(body_rows)}</tbody></table>"


def export_html(question, narrative, df, group_label, metric_label, chart_path, output_dir, filename="rapport.html"):
    os.makedirs(output_dir, exist_ok=True)
    path = os.path.join(output_dir, filename)

    chart_base64 = _image_to_base64(chart_path)
    table_html = _render_table_html(df)
    generated_at = datetime.now().strftime("%d/%m/%Y à %H:%M")

    html = f"""<!DOCTYPE html>
<html lang="fr">
<head>
<meta charset="UTF-8">
<title>Rapport - {question}</title>
<style>
    body {{
        font-family: 'Segoe UI', Calibri, Arial, sans-serif;
        background: #FAFBFD;
        color: #1C1C1C;
        margin: 0;
        padding: 0;
    }}
    .page {{
        max-width: 960px;
        margin: 40px auto;
        background: white;
        border: 1px solid #{ICE};
        border-radius: 10px;
        overflow: hidden;
        box-shadow: 0 8px 24px rgba(30, 39, 97, 0.08);
    }}
    .header {{
        background: #{NAVY};
        color: white;
        padding: 32px 40px;
    }}
    .header .kicker {{
        text-transform: uppercase;
        letter-spacing: 2px;
        font-size: 12px;
        color: #{ICE};
        font-weight: 600;
        margin: 0 0 8px 0;
    }}
    .header h1 {{
        margin: 0;
        font-size: 24px;
        font-weight: 700;
    }}
    .header .meta {{
        margin-top: 10px;
        font-size: 12.5px;
        color: #{ICE};
    }}
    .section {{
        padding: 28px 40px;
    }}
    .section h2 {{
        font-size: 14px;
        text-transform: uppercase;
        letter-spacing: 1px;
        color: #{MIDBLUE};
        margin: 0 0 14px 0;
    }}
    .narrative {{
        background: #{LIGHT_BG};
        border-left: 4px solid #{MIDBLUE};
        padding: 16px 20px;
        border-radius: 6px;
        font-size: 14.5px;
        line-height: 1.6;
        color: #1C1C1C;
    }}
    .chart {{
        text-align: center;
        padding: 10px 40px 28px 40px;
    }}
    .chart img {{
        max-width: 100%;
        border-radius: 6px;
    }}
    table {{
        width: 100%;
        border-collapse: collapse;
        font-size: 13px;
    }}
    thead th {{
        background: #{NAVY};
        color: white;
        text-align: left;
        padding: 10px 12px;
        font-weight: 600;
    }}
    tbody td {{
        padding: 9px 12px;
        border-bottom: 1px solid #{ICE};
    }}
    tbody tr.alt {{
        background: #{LIGHT_BG};
    }}
    td.num {{
        text-align: right;
        font-variant-numeric: tabular-nums;
    }}
    .footer {{
        padding: 18px 40px 32px 40px;
        font-size: 11.5px;
        color: #{MUTED};
        border-top: 1px solid #{ICE};
    }}
</style>
</head>
<body>
<div class="page">
    <div class="header">
        <p class="kicker">Rapport automatique - Agent RAG Factures</p>
        <h1>{metric_label} par {group_label}</h1>
        <div class="meta">Question : "{question}" &nbsp;&middot;&nbsp; Généré le {generated_at}</div>
    </div>
    <div class="section">
        <h2>Synthèse</h2>
        <div class="narrative">{narrative}</div>
    </div>
    <div class="chart">
        <img src="data:image/png;base64,{chart_base64}" alt="Graphique du rapport">
    </div>
    <div class="section">
        <h2>Détail des données</h2>
        {table_html}
    </div>
    <div class="footer">
        Rapport généré automatiquement à partir de la base de facturation (PostgreSQL / Neon).
        Export également disponible au format Excel dans le même dossier.
    </div>
</div>
</body>
</html>
"""

    with open(path, "w", encoding="utf-8") as file:
        file.write(html)

    return path


def generate_report(conn_holder, question, narrative_fn=None, output_dir=None):
    output_dir = output_dir or os.path.join(
        "reports", datetime.now().strftime("%Y%m%d_%H%M%S")
    )

    df, group_label, metric_label = build_report_dataframe(conn_holder, question)

    chart_path = render_chart(df, group_label, metric_label, output_dir)
    excel_path = export_excel(df, output_dir)

    narrative = narrative_fn(df, group_label, metric_label) if narrative_fn else ""
    html_path = export_html(question, narrative, df, group_label, metric_label, chart_path, output_dir)

    return {
        "dataframe": df,
        "group_label": group_label,
        "metric_label": metric_label,
        "excel_path": excel_path,
        "chart_path": chart_path,
        "html_path": html_path,
        "narrative": narrative,
    }